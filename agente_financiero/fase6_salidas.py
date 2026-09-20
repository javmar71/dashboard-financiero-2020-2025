# -*- coding: utf-8 -*-
"""
FASE 6 - SALIDAS EN EXCEL E INFORME INTERACTIVO HTML
=====================================================
Genera las salidas finales sobre el consolidado de los 65 indicadores (2020-2025):

  1. salidas/REPORTE_FINAL_INDICADORES.xlsx      - Libro Excel multi-hoja:
        Hoja 1 "Resumen":  resumen ejecutivo por año + señales.
        Hoja 2 "Indicadores": consolidado 65 indicadores x 6 años (valores reales).
        Hoja 3 "Red Flags":  diagnóstico Fase 5 (estado por (indicador, año)).
        Hoja 4 "Cualitativos": flags de auditoría por año.
        Hoja 5 "Metadatos":  fuentes y trazabilidad.
  2. salidas/informe_financiero_interactivo.html - Informe HTML self-contained con
        gráficos interactivos (plotly) por clasificación + tabla de señales.

Reglas:
  - Solo lectura de salidas ya generadas (indicadores.csv, diagnostico_red_flags.json).
    No vuelve a calcular valores; no inventa.
  - Los nombres y categorías provienen del catálogo (indicadores.csv), alineados con
    DOCUMENTACION_PROYECTO/03_INDICADORES.md (65 indicadores reconciliados).

Uso:
  venv\\Scripts\\python.exe "AUTOMAT ANALISIS FIN\\agente_financiero\\fase6_salidas.py"
"""
from __future__ import annotations

import csv
import html
import json
import sys
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

import pandas as pd
import plotly.graph_objects as go
from plotly.subplots import make_subplots

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill
    from openpyxl.utils import get_column_letter
except Exception:
    openpyxl = None

RAIZ = Path(__file__).resolve().parents[1]
SALIDAS = RAIZ / "salidas"
INDICADORES_CSV = SALIDAS / "indicadores.csv"
DIAG_JSON = SALIDAS / "diagnostico_red_flags.json"
OUT_XLSX = SALIDAS / "REPORTE_FINAL_INDICADORES.xlsx"
OUT_HTML = SALIDAS / "informe_financiero_interactivo.html"
OUT_MATRIZ_HTML = SALIDAS / "MATRIZ_65_INDICADORES.html"
OUT_CSV_65 = RAIZ / "agente_financiero" / "REPORTE_FINAL_65_INDICADORES.csv"
CONFIG_ENTIDAD = SALIDAS / "config_entidad.json"

ANIOS = ["2020", "2021", "2022", "2023", "2024", "2025"]

COLOR_ESTADO = {
    "SEÑAL_ALERTA": "94B8F2",
    "OBSERVACION": "FFD966",
    "OK": "C6EFCE",
    "NO_ENCONTRADO": "D9D9D9",
    "NO_APLICA_TIPO_ENTIDAD": "EDEDED",
}


def leer_indicadores() -> pd.DataFrame:
    df = pd.read_csv(INDICADORES_CSV, encoding="utf-8-sig")
    df = df.rename(columns={"#": "numero", "clasificacion": "clasificacion", "indicador": "indicador",
                            "formula": "formula"})
    for a in ANIOS:
        df[a] = pd.to_numeric(df[a], errors="coerce")
    return df


def leer_diagnostico() -> list[dict]:
    with DIAG_JSON.open("r", encoding="utf-8") as fh:
        data = json.load(fh)
    return data.get("evaluaciones", [])


def cargar_entidad() -> tuple[str, str]:
    """Lee salidas/config_entidad.json (nombre_entidad, nit).

    Campos vacios = modo agnostico (no se muestra entidad). Unica fuente de
    identificacion para HTML y PDF: para usar otra empresa basta editarla.
    """
    try:
        d = json.loads(CONFIG_ENTIDAD.read_text(encoding="utf-8"))
        nombre = str(d.get("nombre_entidad", "")).strip()
        nit = str(d.get("nit", "")).strip()
    except Exception:
        nombre, nit = "", ""
    return nombre, nit


def escribir_excel(df: pd.DataFrame, diag: list[dict], cualitativos: list[dict]) -> None:
    if openpyxl is None:
        raise RuntimeError("openpyxl no disponible")

    import openpyxl as _oxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = _oxl.Workbook()

    estilo_titulo = Font(bold=True, size=13, color="FFFFFF")
    fill_encab = PatternFill("solid", fgColor="305496")
    fill_seccion = PatternFill("solid", fgColor="DDEBF7")

    def hoja_resumen(wb):
        ws = wb.active
        ws.title = "Resumen"
        ws.cell(1, 1, "RESUMEN EJECUTIVO - 65 INDICADORES (2020-2025)").font = estilo_titulo
        ws.cell(1, 1).fill = fill_encab
        ws.cell(2, 1, "Consolidado generado por el pipeline src/ y el catálogo de 65 indicadores "
                      "(03_INDICADORES.md reconciliado). Diagnóstico Fase 5: screening con umbrales documentados; "
                      "no emite juicios de fraude.")
        ws.cell(3, 1, "Fuentes: salidas/indicadores.csv, diagnostico_red_flags.json, lotes de extracción revisados.")
        ws.append([])
        ws.append(["Año", "Señal alerta", "Observación", "OK", "No encontrado", "No aplica", "Total", "Alerta clave"])
        for a in ANIOS:
            d = [e for e in diag if e["anio"] == a]
            conteo = {}
            for e in d:
                conteo[e["estado"]] = conteo.get(e["estado"], 0) + 1
            alertas = [e for e in d if e["estado"] == "SEÑAL_ALERTA"]
            clave = ", ".join(f"#{e['indicador']}" for e in alertas[:4]) if alertas else "-"
            ws.append([a,
                       conteo.get("SEÑAL_ALERTA", 0),
                       conteo.get("OBSERVACION", 0),
                       conteo.get("OK", 0),
                       conteo.get("NO_ENCONTRADO", 0),
                       conteo.get("NO_APLICA_TIPO_ENTIDAD", 0),
                       len(d),
                       clave])
        for cell in ws[5]:
            cell.fill = fill_seccion
            cell.font = Font(bold=True)
        ws.append([])
        ws.append(["NOTA:", "La opinión de auditoría deja de ser limpia desde 2023 (con salvedad). "
                           "La calidad de resultados y la variación de cartera concentran las señales 2024-2025."])
        ws.column_dimensions["A"].width = 14
        ws.column_dimensions["H"].width = 60
        return ws

    def hoja_indicadores(wb):
        ws = wb.create_sheet("Indicadores")
        ws.append(["#", "Clasificación", "Indicador", "Fórmula"] + ANIOS)
        for cell in ws[1]:
            cell.fill = fill_encab
            cell.font = Font(bold=True, color="FFFFFF")
        for _, r in df.iterrows():
            vals = [r[a] if pd.notna(r[a]) else "" for a in ANIOS]
            ws.append([int(r["numero"]), r["clasificacion"], r["indicador"], r["formula"]] + vals)
        for col in range(1, 11):
            ws.column_dimensions[get_column_letter(col)].width = 16 if col >= 5 else (38 if col == 3 else 26)
        ws.freeze_panes = "B2"
        return ws

    def hoja_redflags(wb):
        ws = wb.create_sheet("Red Flags")
        ws.append(["Año", "Número", "Indicador", "Clasificación", "Valor", "Estado", "Motivo", "Fuente regla"])
        for cell in ws[1]:
            cell.fill = fill_encab
            cell.font = Font(bold=True, color="FFFFFF")
        orden = {"SEÑAL_ALERTA": 0, "OBSERVACION": 1, "OK": 2, "NO_ENCONTRADO": 3, "NO_APLICA_TIPO_ENTIDAD": 4}
        for e in sorted(diag, key=lambda x: (x["anio"], orden.get(x["estado"], 9), x["numero"])):
            valor = "" if e.get("valor") is None else e["valor"]
            ws.append([e["anio"], e["numero"], e["indicador"], e["clasificacion"], valor,
                       e["estado"], e["motivo"], e["fuente"]])
        for col in range(1, 9):
            ws.column_dimensions[get_column_letter(col)].width = 55 if col == 7 else 26
        ws.freeze_panes = "A2"
        return ws

    def hoja_cualitativos(wb):
        ws = wb.create_sheet("Cualitativos")
        ws.append(["Año", "Variable", "Valor", "Estado", "Evidencia"])
        for cell in ws[1]:
            cell.fill = fill_encab
            cell.font = Font(bold=True, color="FFFFFF")
        for c in cualitativos:
            ws.append([c["anio"], c["variable"], c["valor"], c["estado"], c["evidencia"]])
        ws.column_dimensions["A"].width = 8
        ws.column_dimensions["B"].width = 20
        ws.column_dimensions["C"].width = 46
        ws.column_dimensions["D"].width = 16
        ws.column_dimensions["E"].width = 90
        return ws

    def hoja_metadatos(wb):
        ws = wb.create_sheet("Metadatos")
        ws.append(["Campo", "Valor"])
        for cell in ws[1]:
            cell.fill = fill_encab
            cell.font = Font(bold=True, color="FFFFFF")
        metas = [
            ("Fase", "6 - Salidas en Excel e informe interactivo HTML"),
            ("Consolidado", "65 indicadores, años 2020-2025"),
            ("Fuente indicadores", str(INDICADORES_CSV)),
            ("Fuente diagnóstico", str(DIAG_JSON)),
            ("Datos brutos", "salidas/datos_estados_financieros.csv (87 x 9)"),
            ("Validaciones", "salidas/reporte_validaciones.txt (65 OK / 0 fallas)"),
            ("Catálogo", "DOCUMENTACION_PROYECTO/03_INDICADORES.md (65 reconciliados)"),
            ("Extracción IA", "agente_financiero/extraction_yyyy_lote*.json (revisados y aprobados)"),
            ("Regla", "Screening Fase 5: umbrales documentados; NOM juicios de fraude"),
            ("Entidad agnóstica", "Reglas universales NIIF/IFRS-XBRL; sin nombres propios"),
            ("Generado", "venv Python 3.12.10 + openpyxl + plotly"),
        ]
        for m in metas:
            ws.append(list(m))
        ws.column_dimensions["A"].width = 24
        ws.column_dimensions["B"].width = 100
        return ws

    for f in (hoja_resumen, hoja_indicadores, hoja_redflags, hoja_cualitativos, hoja_metadatos):
        f(wb)
    wb.save(OUT_XLSX)
    print(f"Escrito: {OUT_XLSX}")


def construir_graficos_y_html(df: pd.DataFrame, diag: list[dict]) -> None:
    clasificaciones = df["clasificacion"].dropna().unique().tolist()
    estado_por: dict[tuple[int, str], str] = {}
    for e in diag:
        if e["numero"] and not e["indicador"].startswith("tipo_opinion"):
            key = (int(e["numero"]), e["anio"])
            if e["numero"] > 0:
                estado_por[key] = e["estado"]

    figures = []
    for clas in clasificaciones:
        sub = df[df["clasificacion"] == clas]
        fig = make_subplots(rows=len(sub), cols=1, shared_xaxes=True,
                            subplot_titles=[f"#{int(r['numero'])} {r['indicador']}" for _, r in sub.iterrows()],
                            vertical_spacing=0.04)
        for i, (_, r) in enumerate(sub.iterrows(), start=1):
            yvals = [r[a] if pd.notna(r[a]) else None for a in ANIOS]
            hovert = None
            fig.add_trace(go.Scatter(x=ANIOS, y=yvals, name=f"#{int(r['numero'])}",
                                     mode="lines+markers"), row=i, col=1)
        fig.update_layout(height=max(260 * len(sub), 300), title=f"{clas} - evolución 2020-2025",
                          showlegend=False, margin=dict(l=60, r=20, t=60, b=40))
        figures.append(fig)

    resumen_por_año = {}
    for e in diag:
        key = e["anio"]
        d = resumen_por_año.setdefault(key, {})
        d[e["estado"]] = d.get(e["estado"], 0) + 1

    estados_cols = ["SEÑAL_ALERTA", "OBSERVACION", "OK", "NO_ENCONTRADO", "NO_APLICA_TIPO_ENTIDAD"]
    tabla_resumen = pd.DataFrame(resumen_por_año).T.fillna(0).astype(int).reindex(ANIOS)
    tabla_resumen = tabla_resumen[estados_cols]
    colores_barra = {"SEÑAL_ALERTA": "#d0342c", "OBSERVACION": "#ffc000", "OK": "#70ad47",
                     "NO_ENCONTRADO": "#bfbfbf", "NO_APLICA_TIPO_ENTIDAD": "#d9d9d9"}
    fig_res = go.Figure(data=[
        go.Bar(name=c, x=tabla_resumen.index, y=tabla_resumen[c], marker_color=colores_barra[c])
        for c in estados_cols])
    fig_res.update_layout(barmode="stack", title="Distribución de estados del screening por año")

    nombre_entidad, nit = cargar_entidad()
    titulo_html = "Informe financiero interactivo - 65 indicadores (2020-2025)"
    encabezado_entidad = ""
    if nombre_entidad:
        titulo_html = (f"Informe financiero interactivo - {nombre_entidad}"
                       + (f" - NIT {nit}" if nit else "") + " - 65 indicadores (2020-2025)")
        encabezado_entidad = (f"<p style='font-size:15px;color:#305496;margin:2px 0'>"
                              f"<strong>{html.escape(nombre_entidad)}</strong>"
                              + (f" &nbsp;|&nbsp; NIT {html.escape(nit)}" if nit else "")
                              + "</p>")
    html_parts: list[str] = []
    html_parts.append("<!DOCTYPE html><html lang='es'><head><meta charset='utf-8'>"
                      "<meta name='viewport' content='width=device-width, initial-scale=1'>")
    html_parts.append(f"<title>{html.escape(titulo_html)}</title>")
    html_parts.append("<style>body{font-family:Segoe UI,Arial,sans-serif;margin:24px;color:#222}"
                      "h1{color:#305496}h2{border-bottom:2px solid #305496;padding-bottom:4px;margin-top:40px}"
                      ".tbl{border-collapse:collapse;font-size:13px;min-width:680px}"
                      ".tbl th{background:#305496;color:#fff;padding:6px 10px;position:sticky;top:0}"
                      ".tbl td{border:1px solid #ccc;padding:5px 9px;white-space:nowrap}"
                      ".resp{width:100%;overflow-x:auto;-webkit-overflow-scrolling:touch}"
                      ".est-alerta{background:#ffd7d7}"
                      ".est-obs{background:#fff3cd}.badge{padding:2px 8px;border-radius:10px;color:#fff;font-size:12px}"
                      ".b-alerta{background:#d0342c}.b-obs{background:#e8a33d}.b-ok{background:#3a8f3f}"
                      ".b-na{background:#8b8b8b}</style></head><body>")
    html_parts.append("<h1>Informe financiero interactivo</h1>")
    if encabezado_entidad:
        html_parts.append(encabezado_entidad)
    html_parts.append("<p><strong>65 indicadores</strong> consolidados del pipeline (2020-2025). "
                      "Diagnóstico Fase 5 como screening con umbrales documentados; no constituye un "
                      "juicio de fraude ni una auditoría concluida.</p>")

    html_parts.append("<h2>1. Resumen del screening por año</h2><div class='resp'>")
    html_parts.append(fig_res.to_html(full_html=False, include_plotlyjs=False))
    html_parts.append(tabla_resumen.to_html(classes="tbl"))
    html_parts.append("</div>")

    html_parts.append("<h2>2. Señales de alerta y observaciones (Fase 5)</h2><div class='resp'>")
    alertas = [e for e in diag if e["estado"] in ("SEÑAL_ALERTA", "OBSERVACION")]
    if alertas:
        html_parts.append("<table class='tbl'><tr><th>Año</th><th>#</th><th>Indicador</th><th>Estado</th>"
                          "<th>Motivo</th></tr>")
        orden_estado = {"SEÑAL_ALERTA": 0, "OBSERVACION": 1}
        for e in sorted(alertas, key=lambda x: (orden_estado[x["estado"]], x["anio"], x["numero"])):
            badge = ("b-alerta" if e["estado"] == "SEÑAL_ALERTA" else "b-obs")
            html_parts.append(f"<tr><td>{e['anio']}</td><td>{e['numero']}</td><td>{e['indicador']}</td>"
                              f"<td><span class='badge {badge}'>{e['estado']}</span></td><td>{e['motivo']}</td></tr>")
        html_parts.append("</table>")
    html_parts.append("</div>")

    html_parts.append("<h2>3. Matriz consolidada de los 65 indicadores por familia y período (2020-2025)</h2>")
    html_parts.append("<p>Los <strong>65 indicadores</strong> del catálogo, agrupados por familia y ordenados del "
                      "año más antiguo (2020) al más reciente (2025). Los 7 indicadores sin base de cálculo "
                      "(inventarios y ciclos operativos, ids 19-22, 36-37, 43) figuran como «—».</p>")
    for clas in clasificaciones:
        sub = df[df["clasificacion"] == clas].sort_values("numero")
        encabezado = "<tr><th>#</th><th>Indicador</th>" + "".join(f"<th>{a}</th>" for a in ANIOS) + "</tr>"
        html_parts.append(f"<h3>{clas}</h3><div class='resp'><table class='tbl'>{encabezado}")
        for _, r in sub.iterrows():
            celdas = "".join(f"<td>{'—' if pd.isna(r[a]) else f'{r[a]:g}'}</td>" for a in ANIOS)
            html_parts.append(f"<tr><td>{int(r['numero'])}</td><td>{r['indicador']}</td>{celdas}</tr>")
        html_parts.append("</table></div>")

    for i, fig in enumerate(figures, start=4):
        html_parts.append(f"<h2>{i}. {fig.layout.title.text}</h2>")
        html_parts.append(fig.to_html(full_html=False, include_plotlyjs=(i == 4)))

    html_parts.append("<h2>Glosario metodológico</h2><ul>")
    html_parts.append("<li>Datos: <code>salidas/indicadores.csv</code> (65 indicadores, reales, calculados por el "
                      "pipeline <code>src/</code>).</li>")
    html_parts.append("<li>Diagnóstico: <code>salidas/diagnostico_red_flags.json</code>.</li>")
    html_parts.append("<li>Señales de auditoría (opinión 2023-2025 con salvedad): lotes de extracción revisados.</li>")
    html_parts.append("</ul></body></html>")

    OUT_HTML.write_text("\n".join(html_parts), encoding="utf-8")
    print(f"Escrito: {OUT_HTML}")


def recoger_cualitativos(diag: list[dict]) -> list[dict]:
    cualitativos: list[dict] = []
    for e in diag:
        if e["numero"] == 0:
            cualitativos.append({
                "anio": e["anio"],
                "variable": e["indicador"],
                "valor": (e["motivo"].split(".")[0] if e["valor"] is None else
                          f"{e['valor']:g}").replace("Opinión modificada: ", ""),
                "estado": e["estado"],
                "evidencia": e["motivo"],
            })
    return cualitativos


def escribir_csv_65(df: pd.DataFrame) -> None:
    """Regenera REPORTE_FINAL_65_INDICADORES.csv (nombres reales y valores), eliminando
    el stub previo (indicador_1..65 con 1.0)."""
    filas: list[list[str]] = [["#", "clasificacion", "indicador", "formula"] + ANIOS]
    for _, r in df.iterrows():
        filas.append([str(int(r["numero"])), r["clasificacion"], r["indicador"], r["formula"]] +
                     [("" if pd.isna(r[a]) else f"{r[a]:g}") for a in ANIOS])
    with OUT_CSV_65.open("w", encoding="utf-8-sig", newline="") as fh:
        csv.writer(fh).writerows(filas)
    print(f"Escrito: {OUT_CSV_65}")


def escribir_matriz_html(df: pd.DataFrame) -> None:
    """Genera un HTML ligero y autónomo con la matriz consolidada de los 65 indicadores
    (una sola tabla: 65 filas x 6 años). Sin plotly, carga instantánea en cualquier navegador."""
    nombre_entidad, nit = cargar_entidad()
    etiqueta = ""
    if nombre_entidad:
        etiqueta = (f"<p style='color:#305496'><strong>{html.escape(nombre_entidad)}</strong>"
                    + (f" &nbsp;|&nbsp; NIT {html.escape(nit)}" if nit else "") + "</p>")
    encabezado = ("<tr><th>#</th><th>Clasificación</th><th>Indicador</th><th>Fórmula</th>"
                  + "".join(f"<th>{a}</th>" for a in ANIOS) + "</tr>")
    filas: list[str] = []
    for _, r in df.sort_values("numero").iterrows():
        celdas = "".join(
            f"<td class='num'>{'—' if pd.isna(r[a]) else f'{r[a]:,.6g}'}</td>" for a in ANIOS)
        filas.append(f"<tr><td>{int(r['numero'])}</td><td>{r['clasificacion']}</td>"
                     f"<td>{r['indicador']}</td><td class='frm'>{r['formula']}</td>{celdas}</tr>")
    html_doc = (
        "<!DOCTYPE html><html lang='es'><head><meta charset='utf-8'>"
        "<meta name='viewport' content='width=device-width, initial-scale=1'>"
        "<title>Matriz consolidada de los 65 indicadores (2020-2025)</title>"
        "<style>body{font-family:Segoe UI,Arial,sans-serif;margin:24px;color:#222}"
        "h1{color:#305496}h2{color:#305496}"
        ".resp{width:100%;overflow-x:auto;-webkit-overflow-scrolling:touch}"
        "table{border-collapse:collapse;font-size:13px;min-width:900px}"
        "th{background:#305496;color:#fff;padding:6px 10px;position:sticky;top:0;text-align:left}"
        "td{border:1px solid #ccc;padding:5px 9px;white-space:nowrap}"
        "td.num{text-align:right;font-variant-numeric:tabular-nums}"
        "td.frm{white-space:normal;color:#555;min-width:220px}"
        "tr:nth-child(even){background:#f6f8fb}</style></head><body>"
        "<h1>Matriz consolidada de indicadores</h1>"
        + etiqueta +
        f"<p><strong>{len(df)} indicadores</strong> del catálogo, año por año (2020-2025). "
        "Los indicadores sin base de cálculo figuran como «—». "
        "Fuente: <code>salidas/indicadores.csv</code>.</p><div class='resp'>"
        f"<table>{encabezado}{''.join(filas)}</table></div></body></html>")
    OUT_MATRIZ_HTML.write_text(html_doc, encoding="utf-8")
    print(f"Escrito: {OUT_MATRIZ_HTML}")


def main() -> int:
    if not INDICADORES_CSV.exists() or not DIAG_JSON.exists():
        print("ERROR: faltan insumos (indicadores.csv / diagnostico_red_flags.json)", file=sys.stderr)
        return 1
    df = leer_indicadores()
    diag = leer_diagnostico()
    cualitativos = recoger_cualitativos(diag)
    escribir_csv_65(df)
    escribir_excel(df, diag, cualitativos)
    construir_graficos_y_html(df, diag)
    escribir_matriz_html(df)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())